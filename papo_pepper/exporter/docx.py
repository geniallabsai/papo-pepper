"""Exporta o paper para .docx profissional: capa, sumário (campo TOC), estilos, tabelas, código."""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

COR_TINTA = RGBColor(0x20, 0x24, 0x2B)
COR_MARCA = RGBColor(0xC8, 0x44, 0x2C)
COR_SUAVE = RGBColor(0x5A, 0x64, 0x72)

_FENCE = re.compile(r"^```\w*\s*$")
_LINHA_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|!\[[^\]]*\]\([^)]+\)|\[[^\]]+\]\([^)]+\)|\*[^*]+\*)")


def gerar_docx(md_caminho: str | Path, out: str | Path, *, titulo: str | None = None,
               autor: str = "Equipe Papo Pepper", data: str | None = None,
               resumo: str = "", palavras_chave: str = "") -> Path:
    md_caminho = Path(md_caminho)
    out = Path(out)
    doc = Document()
    _estilos(doc)
    _capa(doc, titulo or _titulo_do(md_caminho), autor, data or date.today().isoformat(), resumo, palavras_chave)
    _campo(doc, r'TOC \o "1-3" \h \z \u')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("(no Word: clique aqui e pressione F9 para atualizar o sumário)")
    r.font.size = Pt(8)
    r.font.color.rgb = COR_SUAVE
    doc.add_page_break()
    _corpo(doc, md_caminho.read_text(encoding="utf-8"), md_caminho.parent)
    _rodape_pagina(doc)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def _titulo_do(md_caminho: Path) -> str:
    m = re.search(r"^#\s+(.+)$", md_caminho.read_text(encoding="utf-8"), flags=re.M)
    return m.group(1).strip() if m else md_caminho.stem


def _estilos(doc: Document) -> None:
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = COR_TINTA
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for nome, tam, cor in (("Heading 1", 20, COR_MARCA), ("Heading 2", 15.5, COR_TINTA),
                           ("Heading 3", 13, COR_TINTA), ("Heading 4", 11.5, COR_SUAVE)):
        h = st[nome]
        h.font.name = "Calibri"
        h.font.size = Pt(tam)
        h.font.bold = True
        h.font.color.rgb = cor
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)


def _capa(doc: Document, titulo: str, autor: str, data: str, resumo: str, palavras_chave: str) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("P A P O   P E P P E R")
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = COR_MARCA
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(titulo)
    r2.font.size = Pt(28)
    r2.font.bold = True
    r2.font.color.rgb = COR_TINTA
    p3 = doc.add_paragraph()
    pPr = p3._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C8442C")
    pBdr.append(bottom)
    pPr.append(pBdr)
    for rotulo, valor in ((autor, data),):
        pm = doc.add_paragraph()
        pm.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rm = pm.add_run(f"{rotulo}  ·  {valor}")
        rm.font.size = Pt(11)
        rm.font.color.rgb = COR_SUAVE
    if resumo:
        pr = doc.add_paragraph()
        pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pr.add_run("Resumo")
        rr.font.bold = True
        rr.font.size = Pt(11)
        rr.font.color.rgb = COR_MARCA
        pt = doc.add_paragraph()
        pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rt = pt.add_run(resumo[:600])
        rt.font.size = Pt(9.5)
        rt.font.italic = True
    if palavras_chave:
        pk = doc.add_paragraph()
        pk.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rk = pk.add_run(f"Palavras-chave: {palavras_chave}")
        rk.font.size = Pt(9)
        rk.font.color.rgb = COR_SUAVE
    doc.add_page_break()


def _campo(doc: Document, instr: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Sumário — atualize o campo (F9) no Word"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def _sombear(par, cor: str) -> None:
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    pPr.append(shd)


def _bloco_codigo(doc: Document, codigo: str) -> None:
    for linha in codigo.splitlines() or [""]:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        _sombear(par, "F2F2F2")
        run = par.add_run(linha if linha else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x33, 0x3A, 0x45)


def _tabela(doc: Document, linhas_t: list[str]) -> None:
    celulas = [[c.strip() for c in l.strip().strip("|").split("|")] for l in linhas_t if l.strip()]
    if not celulas:
        return
    larguras = max(len(c) for c in celulas)
    t = doc.add_table(rows=len(celulas), cols=larguras)
    try:
        t.style = "Light Grid Accent 1"
    except KeyError:
        pass
    for r_i, linha in enumerate(celulas):
        for c_i in range(larguras):
            texto = linha[c_i] if c_i < len(linha) else ""
            cel = t.cell(r_i, c_i)
            cel.text = ""
            par = cel.paragraphs[0]
            _inline(par, texto)
            if r_i == 0:
                for run in par.runs:
                    run.bold = True


def _inline(par, texto: str) -> None:
    if not texto:
        par.add_run("")
        return
    for parte in _LINHA_INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            par.add_run(parte[2:-2]).bold = True
        elif parte.startswith("`") and parte.endswith("`") and len(parte) > 2:
            r = par.add_run(parte[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            rPr = r._r.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F2F2F2")
            rPr.append(shd)
        elif parte.startswith("!["):
            rotulo = re.sub(r"!\[([^\]]*)\]", r"\1", parte)
            r = par.add_run(f"[figura: {rotulo}]")
            r.italic = True
        elif parte.startswith("[") and "](" in parte:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", parte)
            rotulo, url = m.group(1), m.group(2)
            r = par.add_run(rotulo if rotulo != url else url)
            r.font.color.rgb = RGBColor(0x0F, 0x6A, 0xB4)
            r.underline = True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            par.add_run(parte[1:-1]).italic = True
        else:
            par.add_run(parte)


def _corpo(doc: Document, md: str, base_dir: Path) -> None:
    linhas = md.splitlines()
    i, n = 0, len(linhas)
    while i < n:
        linha = linhas[i]
        if not linha.strip():
            i += 1
            continue
        if _FENCE.match(linha.strip()):
            i += 1
            buf = []
            while i < n and not _FENCE.match(linhas[i].strip()):
                buf.append(linhas[i])
                i += 1
            i += 1
            _bloco_codigo(doc, "\n".join(buf))
            continue
        if linha.lstrip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:\-|]+\|\s*$", linhas[i + 1]):
            tabela = [linha]
            while i < n and linhas[i].lstrip().startswith("|"):
                tabela.append(linhas[i])
                i += 1
            _tabela(doc, [l for j, l in enumerate(tabela) if j != 1])
            continue
        hm = re.match(r"^(#{1,6})\s+(.*)$", linha)
        if hm:
            texto = re.sub(r"<[^>]+>", "", hm.group(2)).strip()
            doc.add_heading(texto, level=min(len(hm.group(1)), 4))
            i += 1
            continue
        if linha.strip() in ("---", "***", "___"):
            i += 1
            continue
        if re.match(r"^\s*[-*+]\s+", linha):
            while i < n and re.match(r"^\s*[-*+]\s+", linhas[i]):
                conteudo = re.sub(r"^\s*[-*+]\s+", "", linhas[i])
                indent = (len(linhas[i]) - len(linhas[i].lstrip())) // 2
                par = doc.add_paragraph(style="List Bullet" if indent == 0 else "List Bullet 2")
                _inline(par, conteudo)
                i += 1
            continue
        if re.match(r"^\s*\d+[.)]\s+", linha):
            while i < n and re.match(r"^\s*\d+[.)]\s+", linhas[i]):
                conteudo = re.sub(r"^\s*\d+[.)]\s+", "", linhas[i])
                par = doc.add_paragraph(style="List Number")
                _inline(par, conteudo)
                i += 1
            continue
        if linha.lstrip().startswith(">"):
            while i < n and (linhas[i].lstrip().startswith(">") or (linhas[i].strip() and i > 0 and linhas[i - 1].lstrip().startswith(">"))):
                conteudo = re.sub(r"^\s*>\s?", "", linhas[i])
                try:
                    par = doc.add_paragraph(style="Intense Quote")
                except KeyError:
                    par = doc.add_paragraph()
                    par.paragraph_format.left_indent = Inches(0.3)
                _inline(par, conteudo)
                i += 1
            continue
        if linha.lstrip().startswith("!["):
            mimg = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", linha.strip())
            if mimg:
                rotulo, src = mimg.group(1), mimg.group(2)
                caminho = Path(src)
                if not caminho.is_absolute():
                    caminho = base_dir / src
                if caminho.exists():
                    try:
                        doc.add_picture(str(caminho), width=Inches(5.9))
                        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        if rotulo:
                            cap = doc.add_paragraph()
                            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            rc = cap.add_run(rotulo)
                            rc.font.size = Pt(9)
                            rc.font.italic = True
                            rc.font.color.rgb = COR_SUAVE
                    except Exception:  # noqa: BLE001 — imagem inválida não derruba o docx
                        pass
                i += 1
                continue
        par = doc.add_paragraph()
        _inline(par, linha)
        i += 1


def _rodape_pagina(doc: Document) -> None:
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p.add_run("Papo Pepper · página ")
        r0.font.size = Pt(8)
        r0.font.color.rgb = COR_SUAVE
        run = p.add_run()
        run.font.size = Pt(8)
        run.font.color.rgb = COR_SUAVE
        fld1 = OxmlElement("w:fldChar")
        fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.text = "PAGE"
        fld2 = OxmlElement("w:fldChar")
        fld2.set(qn("w:fldCharType"), "end")
        run._r.append(fld1)
        run._r.append(instr)
        run._r.append(fld2)
