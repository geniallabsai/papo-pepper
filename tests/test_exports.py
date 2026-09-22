import tempfile
from pathlib import Path

MD = """# Teste de Paper Papo Pepper

**Negrito** e *itálico* e `codigo` e [link](https://example.com).

## Seção um
Parágrafo com equação $E = mc^2$ e referência [1].

| Coluna A | Coluna B |
| --- | --- |
| 1 | dois |

- item um
- item dois

1. passo
2. passo

> citação importante

```python
print("olá")
```
"""


def test_html_contem_sumario_e_conteudo():
    from papo_pepper.exporter import html as html_exp
    h = html_exp.gerar_html(MD, titulo="Teste PP", autor="QA", data="2026-09-22")
    assert "<h1" in h
    assert "Sumário" in h
    assert "Coluna A" in h
    assert "capa" in h


def test_fallback_mini_md():
    from papo_pepper.exporter import html as html_exp
    h = html_exp._mini_md(MD)
    assert "<h1>" in h
    assert "<table" in h


def test_docx_abre_e_tem_paragrafos():
    from docx import Document
    from papo_pepper.exporter import docx as docx_exp
    with tempfile.TemporaryDirectory() as tmp:
        md = Path(tmp) / "p.md"
        md.write_text(MD, encoding="utf-8")
        out = docx_exp.gerar_docx(md, Path(tmp) / "p.docx", titulo="Teste PP", autor="QA")
        assert out.stat().st_size > 5000
        d = Document(str(out))
        assert len(d.paragraphs) > 10
        assert any("Seção um" in p.text for p in d.paragraphs)


def test_pdf_tem_cabecalho_valido():
    from papo_pepper.exporter import pdf as pdf_exp
    with tempfile.TemporaryDirectory() as tmp:
        md = Path(tmp) / "p.md"
        md.write_text(MD, encoding="utf-8")
        out = Path(tmp) / "p.pdf"
        engine = pdf_exp.gerar_pdf(md, out, titulo="Teste PP", autor="QA")
        dados = out.read_bytes()
        assert dados.startswith(b"%PDF"), f"motor {engine} nao gerou PDF valido"
        assert len(dados) > 1000


def test_bibtex_gera_entrada_misc():
    from papo_pepper.exporter import bib
    fontes = [{"id": 1, "titulo": "Um Estudo X", "url": "https://arxiv.org/abs/2025.12345",
               "data": "2025-03-01", "autores": ["Silva"]}]
    with tempfile.TemporaryDirectory() as tmp:
        out = bib.gerar_bibtex(fontes, Path(tmp) / "r.bib")
        texto = out.read_text(encoding="utf-8")
        assert "@misc{" in texto
        assert "arxiv.org" in texto
        assert "Silva" in texto
